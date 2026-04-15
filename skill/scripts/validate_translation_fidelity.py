#!/usr/bin/env python3
"""Check translation fidelity for explanatory drift, omissions, and bad segment granularity."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path


SUSPICIOUS_PATTERNS = [
    "\u4f5c\u8005\u6307\u51fa",
    "\u4f5c\u8005\u8ba4\u4e3a",
    "\u4f5c\u8005\u56e2\u961f",
    "\u4f5c\u8005\u8fdb\u4e00\u6b65",
    "\u672c\u6587\u63d0\u51fa",
    "\u672c\u6587\u8fd8",
    "\u6587\u4e2d\u6307\u51fa",
    "\u8be5\u7814\u7a76\u8868\u660e",
    "\u7814\u7a76\u8005\u53d1\u73b0",
]
REGISTER_DRIFT_PATTERNS = [
    "说白了",
    "简单来说",
    "通俗地说",
    "打个比方",
    "大家知道",
    "总的来说",
    "总之",
    "其实",
    "相当于",
    "就像",
]

SOURCE_SENTENCE_BOUNDARY_RE = re.compile(r"[.!?](?=(?:\s+|$))")
TARGET_SENTENCE_RE = re.compile(r"(?<=[。！？；;!?])")
FIGURE_REF_RE = re.compile(r"\b(?:Fig\.?|Figure)\s*\d+[a-z]?", re.IGNORECASE)
EQUATION_REF_RE = re.compile(r"\bEq\.?\s*\(?\d+\)?", re.IGNORECASE)
SUPPLEMENTARY_REF_RE = re.compile(
    r"\bSupplementary\s+(?:Fig(?:ure)?\.?|Table|Movie)\s*\d+[a-z]?",
    re.IGNORECASE,
)
PARAMETER_SYMBOL_RE = re.compile(
    r"\b(P|OD|DC|SNR|PSNR|NMSE|Correlation|VDR|f\s*0)\s*=",
    re.IGNORECASE,
)
PERCENT_TOKEN_RE = re.compile(r"\b\d+(?:\.\d+)?%")
POWER_TOKEN_RE = re.compile(r"\b10\^\d+\b", re.IGNORECASE)
UNIT_TOKEN_RE = re.compile(r"\b\d+(?:\.\d+)?\s*(MPa|MHz|mm|W)\b", re.IGNORECASE)
COEFFICIENT_TOKEN_RE = re.compile(r"\b([A-Za-z][A-Za-z0-9_]*)\s*=\s*[-+0-9]")
COEFFICIENT_VALUE_RE = re.compile(
    r"\b[A-Za-z][A-Za-z0-9_]*\s*=\s*[-+]?\s*(\d+(?:\.\d+)?)\s*(?:[×x]\s*)?(10\^\d+)?",
    re.IGNORECASE,
)


def load_docx_paragraphs(path: Path) -> list[str]:
    from docx import Document  # type: ignore

    doc = Document(str(path))
    return [p.text.strip() for p in doc.paragraphs if p.text.strip()]


def load_text_paragraphs(path: Path) -> list[str]:
    text = path.read_text(encoding="utf-8-sig")
    return [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]


def load_workbook_segments(path: Path) -> list[dict]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    return payload.get("segments", [])


def body_like_paragraphs(paragraphs: list[str]) -> list[str]:
    filtered = []
    for text in paragraphs:
        if text.startswith("图") or text.lower().startswith("fig."):
            continue
        if re.match(r"^\d+\.\s", text):
            continue
        filtered.append(text)
    return filtered


def count_source_sentences(text: str) -> int:
    normalized = " ".join(text.split())
    if not normalized:
        return 0
    count = len(SOURCE_SENTENCE_BOUNDARY_RE.findall(normalized))
    return max(1, count)


def count_target_sentences(text: str) -> int:
    normalized = text.strip()
    if not normalized:
        return 0
    count = len([part for part in TARGET_SENTENCE_RE.split(normalized) if part.strip()])
    return max(1, count)


def detect_suspicious_voice(paragraphs: list[str]) -> list[dict]:
    hits = []
    for para in paragraphs:
        patterns = [pattern for pattern in SUSPICIOUS_PATTERNS if pattern in para]
        if patterns:
            hits.append({"text": para[:150], "patterns": patterns})
    return hits


def detect_register_drift(paragraphs: list[str]) -> list[dict]:
    hits = []
    for para in paragraphs:
        matched = [pattern for pattern in REGISTER_DRIFT_PATTERNS if pattern in para]
        if "！" in para or "?" in para or "？" in para:
            matched.append("rhetorical_punctuation")
        if matched:
            hits.append({"text": para[:150], "patterns": matched})
    return hits


def normalize_markers(markers: set[str]) -> set[str]:
    return {re.sub(r"\s+", "", marker).lower() for marker in markers}


def extract_structural_markers(text: str) -> dict[str, set[str]]:
    return {
        "figure_refs": normalize_markers(set(FIGURE_REF_RE.findall(text))),
        "equation_refs": normalize_markers(set(EQUATION_REF_RE.findall(text))),
        "supplementary_refs": normalize_markers(set(SUPPLEMENTARY_REF_RE.findall(text))),
        "parameter_symbols": normalize_markers({match.group(1) for match in PARAMETER_SYMBOL_RE.finditer(text)}),
    }


def normalize_text_for_token_match(text: str) -> str:
    return re.sub(r"\s+", "", (text or "")).lower()


def extract_critical_tokens(text: str) -> list[str]:
    tokens: set[str] = set()
    for token in PERCENT_TOKEN_RE.findall(text or ""):
        tokens.add(token.lower())
    for match in UNIT_TOKEN_RE.finditer(text or ""):
        tokens.add(match.group(1).lower())
        number = re.search(r"\d+(?:\.\d+)?", match.group(0))
        if number:
            tokens.add(number.group(0).lower())
    for token in COEFFICIENT_TOKEN_RE.findall(text or ""):
        normalized = token.lower()
        if any(char.isdigit() for char in normalized):
            tokens.add(normalized)
    for value, power in COEFFICIENT_VALUE_RE.findall(text or ""):
        tokens.add(value.lower())
        if power:
            tokens.add(power.lower())
    return sorted(tokens)


def analyze_segments(payload_segments: list[dict]) -> dict:
    source_body_segments = [seg for seg in payload_segments if seg.get("type") == "body"]
    oversized_source_segments = []
    suspicious_sentence_drop = []
    suspicious_length_drop = []
    translated_body_segments = []
    missing_structural_markers = []
    missing_caption_parameters = []
    missing_critical_tokens = []

    for segment in payload_segments:
        source_text = (segment.get("source_text") or "").strip()
        translated_text = (segment.get("translated_text") or "").strip()
        segment_type = segment.get("type")

        if translated_text:
            source_markers = extract_structural_markers(source_text)
            target_markers = extract_structural_markers(translated_text)

            missing_markers = {}
            for key in ("figure_refs", "equation_refs", "supplementary_refs"):
                missing = sorted(source_markers[key] - target_markers[key])
                if missing:
                    missing_markers[key] = missing

            if missing_markers:
                missing_structural_markers.append(
                    {
                        "source_id": segment.get("source_id"),
                        "page_number": segment.get("page_number"),
                        "segment_type": segment_type,
                        "missing_markers": missing_markers,
                        "source_text_preview": source_text[:180],
                        "translated_text_preview": translated_text[:180],
                    }
                )

            if segment_type == "caption":
                missing_parameters = sorted(
                    source_markers["parameter_symbols"] - target_markers["parameter_symbols"]
                )
                if missing_parameters:
                    missing_caption_parameters.append(
                        {
                            "source_id": segment.get("source_id"),
                            "page_number": segment.get("page_number"),
                            "missing_parameter_symbols": missing_parameters,
                            "source_text_preview": source_text[:180],
                            "translated_text_preview": translated_text[:180],
                        }
                    )

            required_tokens = extract_critical_tokens(source_text)
            if required_tokens:
                normalized_target = normalize_text_for_token_match(translated_text)
                missing_tokens = [token for token in required_tokens if token not in normalized_target]
                if missing_tokens:
                    missing_critical_tokens.append(
                        {
                            "source_id": segment.get("source_id"),
                            "page_number": segment.get("page_number"),
                            "missing_tokens": missing_tokens,
                        }
                    )

        if segment_type != "body":
            continue

        source_sentences = count_source_sentences(source_text)
        source_words = len(re.findall(r"[A-Za-z0-9]+", source_text))

        if source_words >= 170 or source_sentences >= 7:
            oversized_source_segments.append(
                {
                    "source_id": segment.get("source_id"),
                    "page_number": segment.get("page_number"),
                    "source_word_count": source_words,
                    "source_sentence_count": source_sentences,
                    "source_text_preview": source_text[:180],
                }
            )

        if translated_text:
            translated_body_segments.append(translated_text)
            target_sentences = count_target_sentences(translated_text)
            target_character_count = len(re.sub(r"\s+", "", translated_text))
            if source_sentences >= 4 and target_sentences <= max(1, source_sentences // 2):
                suspicious_sentence_drop.append(
                    {
                        "source_id": segment.get("source_id"),
                        "page_number": segment.get("page_number"),
                        "source_sentence_count": source_sentences,
                        "target_sentence_count": target_sentences,
                        "source_text_preview": source_text[:180],
                        "translated_text_preview": translated_text[:180],
                    }
                )
            if source_words >= 45 and target_character_count < max(18, int(source_words * 0.35)):
                suspicious_length_drop.append(
                    {
                        "source_id": segment.get("source_id"),
                        "page_number": segment.get("page_number"),
                        "source_word_count": source_words,
                        "target_character_count": target_character_count,
                    }
                )

    return {
        "source_body_segment_count": len(source_body_segments),
        "body_like_paragraph_count": len(translated_body_segments),
        "oversized_source_segments": oversized_source_segments,
        "suspicious_sentence_drop": suspicious_sentence_drop,
        "suspicious_length_drop": suspicious_length_drop,
        "missing_structural_markers": missing_structural_markers,
        "missing_caption_parameters": missing_caption_parameters,
        "missing_critical_tokens": missing_critical_tokens,
        "suspicious_voice_hits": detect_suspicious_voice(translated_body_segments),
        "register_drift_hits": detect_register_drift(translated_body_segments),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("translation", help="Translated DOCX, UTF-8 text, or translation workbook JSON")
    parser.add_argument("--segments", help="Optional source segments JSON from prepare_translation_segments.py")
    args = parser.parse_args()

    translation_path = Path(args.translation).expanduser().resolve()
    report = {"translation": str(translation_path)}

    if translation_path.suffix.lower() == ".docx":
        paragraphs = load_docx_paragraphs(translation_path)
        body_paras = body_like_paragraphs(paragraphs)
        report["paragraph_count"] = len(paragraphs)
        report["body_like_paragraph_count"] = len(body_paras)
        report["suspicious_voice_hits"] = detect_suspicious_voice(body_paras)
        report["register_drift_hits"] = detect_register_drift(body_paras)
    elif translation_path.suffix.lower() == ".json":
        workbook_segments = load_workbook_segments(translation_path)
        report["segment_count"] = len(workbook_segments)
        report.update(analyze_segments(workbook_segments))
    else:
        paragraphs = load_text_paragraphs(translation_path)
        body_paras = body_like_paragraphs(paragraphs)
        report["paragraph_count"] = len(paragraphs)
        report["body_like_paragraph_count"] = len(body_paras)
        report["suspicious_voice_hits"] = detect_suspicious_voice(body_paras)
        report["register_drift_hits"] = detect_register_drift(body_paras)

    if args.segments:
        segments_path = Path(args.segments).expanduser().resolve()
        payload = json.loads(segments_path.read_text(encoding="utf-8"))
        source_body_count = sum(1 for seg in payload.get("segments", []) if seg.get("type") == "body")
        report["prepared_source_body_segment_count"] = source_body_count
        if "body_like_paragraph_count" in report:
            report["body_paragraph_delta"] = report["body_like_paragraph_count"] - source_body_count

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))

    blocking_issues = []
    if report.get("suspicious_voice_hits"):
        blocking_issues.append("suspicious_voice_hits")
    if report.get("register_drift_hits"):
        blocking_issues.append("register_drift_hits")
    if report.get("oversized_source_segments"):
        blocking_issues.append("oversized_source_segments")
    if report.get("suspicious_sentence_drop"):
        blocking_issues.append("suspicious_sentence_drop")
    if report.get("suspicious_length_drop"):
        blocking_issues.append("suspicious_length_drop")
    if report.get("missing_structural_markers"):
        blocking_issues.append("missing_structural_markers")
    if report.get("missing_caption_parameters"):
        blocking_issues.append("missing_caption_parameters")
    if report.get("missing_critical_tokens"):
        blocking_issues.append("missing_critical_tokens")
    if isinstance(report.get("body_paragraph_delta"), int) and report["body_paragraph_delta"] < -3:
        blocking_issues.append("body_paragraph_delta")

    return 2 if blocking_issues else 0


if __name__ == "__main__":
    raise SystemExit(main())
