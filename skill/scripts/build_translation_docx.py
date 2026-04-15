#!/usr/bin/env python3
"""Build a styled DOCX translation deliverable from structured JSON content."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path


REFERENCE_LABEL_RE = re.compile(r"^\s*(?:\[(\d+)\]|(\d+)[\.\)])\s*")
CITATION_RE = re.compile(r"\[(\d+(?:\s*[,;]\s*\d+)*)\]")
OMML_NAMESPACE = (
    'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
)


def load_docx():
    try:
        from docx import Document  # type: ignore
        from docx.enum.style import WD_STYLE_TYPE  # type: ignore
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING  # type: ignore
        from docx.oxml import OxmlElement, parse_xml  # type: ignore
        from docx.oxml.ns import qn  # type: ignore
        from docx.shared import Inches, Pt  # type: ignore
    except ImportError as exc:  # pragma: no cover - dependency guidance
        raise SystemExit(
            "python-docx is required. Install it with `python -m pip install python-docx`."
        ) from exc
    return (
        Document,
        WD_STYLE_TYPE,
        WD_ALIGN_PARAGRAPH,
        WD_LINE_SPACING,
        OxmlElement,
        parse_xml,
        qn,
        Inches,
        Pt,
    )


def apply_font(style, qn, pt, *, ascii_font="Times New Roman", east_asia_font="\u5b8b\u4f53", bold=False):
    style.font.name = ascii_font
    style.font.size = pt
    style.font.bold = bold
    style.element.rPr.rFonts.set(qn("w:ascii"), ascii_font)
    style.element.rPr.rFonts.set(qn("w:hAnsi"), ascii_font)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), east_asia_font)


def ensure_paragraph_style(document, WD_STYLE_TYPE, name: str):
    styles = document.styles
    try:
        return styles[name]
    except KeyError:
        return styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def configure_styles(document, WD_STYLE_TYPE, WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, qn, Pt):
    title_style = document.styles["Title"]
    apply_font(title_style, qn, Pt(16), east_asia_font="\u9ed1\u4f53", bold=True)
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.space_after = Pt(10)

    heading1 = document.styles["Heading 1"]
    apply_font(heading1, qn, Pt(14), east_asia_font="\u9ed1\u4f53", bold=True)
    heading1.paragraph_format.space_before = Pt(10)
    heading1.paragraph_format.space_after = Pt(6)

    heading2 = document.styles["Heading 2"]
    apply_font(heading2, qn, Pt(12), east_asia_font="\u9ed1\u4f53", bold=True)
    heading2.paragraph_format.space_before = Pt(8)
    heading2.paragraph_format.space_after = Pt(4)

    caption = document.styles["Caption"]
    apply_font(caption, qn, Pt(10.5), east_asia_font="\u6977\u4f53")
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    body = ensure_paragraph_style(document, WD_STYLE_TYPE, "PaperBody")
    apply_font(body, qn, Pt(12), east_asia_font="\u5b8b\u4f53")
    body.paragraph_format.first_line_indent = Pt(24)
    body.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    body.paragraph_format.space_after = Pt(4)

    authors = ensure_paragraph_style(document, WD_STYLE_TYPE, "PaperAuthors")
    apply_font(authors, qn, Pt(11), east_asia_font="\u6977\u4f53")
    authors.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    authors.paragraph_format.space_after = Pt(4)

    affiliations = ensure_paragraph_style(document, WD_STYLE_TYPE, "PaperAffiliations")
    apply_font(affiliations, qn, Pt(10.5), east_asia_font="\u5b8b\u4f53")
    affiliations.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    affiliations.paragraph_format.space_after = Pt(4)

    metadata = ensure_paragraph_style(document, WD_STYLE_TYPE, "PaperMetadata")
    apply_font(metadata, qn, Pt(10.5), east_asia_font="\u5b8b\u4f53")
    metadata.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_after = Pt(6)

    reference = ensure_paragraph_style(document, WD_STYLE_TYPE, "PaperReference")
    apply_font(reference, qn, Pt(10.5), east_asia_font="\u5b8b\u4f53")
    reference.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    reference.paragraph_format.left_indent = Pt(21)
    reference.paragraph_format.first_line_indent = Pt(-21)
    reference.paragraph_format.space_after = Pt(2)

    abstract_body = ensure_paragraph_style(document, WD_STYLE_TYPE, "PaperAbstract")
    apply_font(abstract_body, qn, Pt(11), east_asia_font="\u5b8b\u4f53")
    abstract_body.paragraph_format.first_line_indent = Pt(24)
    abstract_body.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    abstract_body.paragraph_format.space_after = Pt(4)

    keyword_style = ensure_paragraph_style(document, WD_STYLE_TYPE, "PaperKeyword")
    apply_font(keyword_style, qn, Pt(10.5), east_asia_font="\u5b8b\u4f53")
    keyword_style.paragraph_format.space_after = Pt(4)

    equation = ensure_paragraph_style(document, WD_STYLE_TYPE, "PaperEquation")
    apply_font(equation, qn, Pt(11), east_asia_font="Cambria Math")
    equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.space_after = Pt(4)

    note = ensure_paragraph_style(document, WD_STYLE_TYPE, "TranslatorNote")
    apply_font(note, qn, Pt(10.5), east_asia_font="\u6977\u4f53")
    note.paragraph_format.space_after = Pt(4)


def extract_reference_label(text: str) -> str | None:
    match = REFERENCE_LABEL_RE.match(text)
    if not match:
        return None
    return match.group(1) or match.group(2)


def segment_bookmark_name(source_id: str) -> str:
    normalized = re.sub(r"[^A-Za-z0-9_-]+", "-", (source_id or "").strip()).strip("-").lower()
    if not normalized:
        raise ValueError("source_id is required for segment bookmarks")
    if normalized[0].isdigit():
        normalized = f"s-{normalized}"
    return f"seg-{normalized}"


def collect_reference_targets(payload: dict) -> dict[str, str]:
    targets: dict[str, str] = {}
    blocks = payload.get("blocks")
    segments = payload.get("segments")

    if isinstance(blocks, list):
        for block in blocks:
            if block.get("type") != "reference":
                continue
            label = extract_reference_label(str(block.get("text", "")))
            if label:
                targets[label] = f"ref-{label}"

    if isinstance(segments, list):
        for segment in segments:
            if segment.get("type") != "reference":
                continue
            label = extract_reference_label(str(segment.get("translated_text") or segment.get("source_text") or ""))
            if label:
                targets[label] = f"ref-{label}"

    return targets


def add_plain_run(paragraph, text: str):
    if text:
        paragraph.add_run(text)


def add_internal_hyperlink(paragraph, text: str, anchor: str, OxmlElement, qn):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")

    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rstyle = OxmlElement("w:rStyle")
    rstyle.set(qn("w:val"), "Hyperlink")
    rpr.append(rstyle)
    run.append(rpr)

    text_element = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" "):
        text_element.set(qn("xml:space"), "preserve")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def append_text_with_reference_links(paragraph, text: str, reference_targets: dict[str, str], OxmlElement, qn):
    cursor = 0
    for match in CITATION_RE.finditer(text):
        add_plain_run(paragraph, text[cursor:match.start()])
        content = match.group(1)
        parts = re.split(r"(\s*[,;]\s*)", content)
        paragraph.add_run("[")
        for part in parts:
            cleaned = part.strip()
            if not cleaned:
                paragraph.add_run(part)
                continue
            if cleaned.isdigit() and cleaned in reference_targets:
                add_internal_hyperlink(paragraph, cleaned, reference_targets[cleaned], OxmlElement, qn)
            else:
                paragraph.add_run(part)
        paragraph.add_run("]")
        cursor = match.end()
    add_plain_run(paragraph, text[cursor:])


def add_bookmark(paragraph, bookmark_name: str, bookmark_state: dict[str, int], OxmlElement, qn):
    bookmark_id = bookmark_state["next_id"]
    bookmark_state["next_id"] += 1

    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), bookmark_name)

    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))

    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_text_paragraph(
    document,
    text: str,
    style_name: str,
    reference_targets: dict[str, str],
    bookmark_state: dict[str, int],
    OxmlElement,
    qn,
    bookmark_name: str | None = None,
):
    paragraph = document.add_paragraph(style=style_name)
    append_text_with_reference_links(paragraph, text, reference_targets, OxmlElement, qn)
    if bookmark_name:
        add_bookmark(paragraph, bookmark_name, bookmark_state, OxmlElement, qn)
    return paragraph


def ensure_omml_namespaces(omml: str) -> str:
    if "xmlns:m=" in omml:
        return omml
    if omml.startswith("<m:oMathPara"):
        return omml.replace("<m:oMathPara", f"<m:oMathPara {OMML_NAMESPACE}", 1)
    if omml.startswith("<m:oMath"):
        return omml.replace("<m:oMath", f"<m:oMath {OMML_NAMESPACE}", 1)
    return omml


def add_equation_paragraph(document, omml: str, parse_xml, bookmark_name: str | None, bookmark_state, OxmlElement, qn):
    paragraph = document.add_paragraph(style="PaperEquation")
    equation_xml = ensure_omml_namespaces(omml.strip())
    paragraph._p.append(parse_xml(equation_xml))
    if bookmark_name:
        add_bookmark(paragraph, bookmark_name, bookmark_state, OxmlElement, qn)
    return paragraph


def add_block(
    document,
    inches,
    WD_ALIGN_PARAGRAPH,
    block: dict,
    base_dir: Path,
    reference_targets: dict[str, str],
    bookmark_state: dict[str, int],
    OxmlElement,
    parse_xml,
    qn,
) -> None:
    block_type = block.get("type")
    text = block.get("text", "")
    bookmark_name = None
    if block.get("source_id"):
        bookmark_name = segment_bookmark_name(str(block["source_id"]))

    if block_type == "title":
        document.add_paragraph(text, style="Title")
    elif block_type == "heading":
        level = max(1, min(int(block.get("level", 1)), 9))
        document.add_heading(text, level=level)
    elif block_type == "authors":
        add_text_paragraph(
            document, text, "PaperAuthors", reference_targets, bookmark_state, OxmlElement, qn, bookmark_name
        )
    elif block_type in {"affiliation", "affiliations"}:
        add_text_paragraph(
            document,
            text,
            "PaperAffiliations",
            reference_targets,
            bookmark_state,
            OxmlElement,
            qn,
            bookmark_name,
        )
    elif block_type in {"doi", "metadata"}:
        add_text_paragraph(
            document, text, "PaperMetadata", reference_targets, bookmark_state, OxmlElement, qn, bookmark_name
        )
    elif block_type == "abstract":
        add_text_paragraph(
            document, text, "PaperAbstract", reference_targets, bookmark_state, OxmlElement, qn, bookmark_name
        )
    elif block_type == "keyword":
        add_text_paragraph(
            document, text, "PaperKeyword", reference_targets, bookmark_state, OxmlElement, qn, bookmark_name
        )
    elif block_type == "paragraph":
        add_text_paragraph(
            document, text, "PaperBody", reference_targets, bookmark_state, OxmlElement, qn, bookmark_name
        )
    elif block_type == "reference":
        label = extract_reference_label(text)
        reference_bookmark = reference_targets.get(label) if label else None
        add_text_paragraph(
            document,
            text,
            "PaperReference",
            reference_targets,
            bookmark_state,
            OxmlElement,
            qn,
            bookmark_name=reference_bookmark or bookmark_name,
        )
    elif block_type == "caption":
        add_text_paragraph(document, text, "Caption", reference_targets, bookmark_state, OxmlElement, qn, bookmark_name)
    elif block_type == "equation":
        equation_omml = (block.get("equation_omml") or block.get("omml") or "").strip()
        if not equation_omml:
            raise SystemExit("Equation blocks require machine-readable OMML. Equation images are not permitted.")
        add_equation_paragraph(document, equation_omml, parse_xml, bookmark_name, bookmark_state, OxmlElement, qn)
    elif block_type == "translator-note":
        paragraph = document.add_paragraph(style="TranslatorNote")
        run = paragraph.add_run("Translator note: ")
        run.bold = True
        paragraph.add_run(text)
        if bookmark_name:
            add_bookmark(paragraph, bookmark_name, bookmark_state, OxmlElement, qn)
    elif block_type == "image":
        if block.get("role") == "equation" or block.get("kind") == "equation":
            raise SystemExit("Equation images are not permitted. Supply editable OMML instead.")
        image_path = (
            (base_dir / block["path"]).resolve()
            if not Path(block["path"]).is_absolute()
            else Path(block["path"])
        )
        width_inches = float(block.get("width_inches", 5.8))
        picture_paragraph = document.add_paragraph()
        picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_paragraph.add_run().add_picture(str(image_path), width=inches(width_inches))
        if bookmark_name:
            add_bookmark(picture_paragraph, bookmark_name, bookmark_state, OxmlElement, qn)
        if block.get("caption"):
            add_text_paragraph(
                document,
                block["caption"],
                "Caption",
                reference_targets,
                bookmark_state,
                OxmlElement,
                qn,
                bookmark_name,
            )
    else:
        if text:
            add_text_paragraph(
                document, text, "PaperBody", reference_targets, bookmark_state, OxmlElement, qn, bookmark_name
            )


def validate_workbook_segments(segments: list[dict]) -> None:
    seen = set()
    missing = []
    duplicates = []
    empty = []
    omitted = []
    equation_missing_omml = []

    for segment in segments:
        source_id = segment.get("source_id", "")
        translated_text = (segment.get("translated_text") or "").strip()
        status = segment.get("status", "")

        if not source_id:
            missing.append(source_id)
        elif source_id in seen:
            duplicates.append(source_id)
        seen.add(source_id)

        if status == "omitted":
            omitted.append(source_id)
        if status != "omitted" and not translated_text and segment.get("type") != "equation":
            empty.append(source_id)
        if segment.get("type") == "equation" and status != "omitted":
            if not (segment.get("equation_omml") or "").strip():
                equation_missing_omml.append(source_id)

    errors = []
    if missing:
        errors.append(f"missing source_id entries: {len(missing)}")
    if duplicates:
        errors.append(f"duplicate source_id entries: {len(duplicates)}")
    if omitted:
        preview = ", ".join(omitted[:10])
        errors.append(f"omitted segments are not permitted: {len(omitted)} ({preview})")
    if empty:
        preview = ", ".join(empty[:10])
        errors.append(f"empty translated_text entries: {len(empty)} ({preview})")
    if equation_missing_omml:
        preview = ", ".join(equation_missing_omml[:10])
        errors.append(
            f"equation segments missing machine-readable OMML: {len(equation_missing_omml)} ({preview})"
        )

    if errors:
        raise SystemExit("Refusing to build DOCX from incomplete workbook: " + "; ".join(errors))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_json", help="Structured translation JSON")
    parser.add_argument("output_docx", help="Output DOCX path")
    args = parser.parse_args()

    (
        Document,
        WD_STYLE_TYPE,
        WD_ALIGN_PARAGRAPH,
        WD_LINE_SPACING,
        OxmlElement,
        parse_xml,
        qn,
        Inches,
        Pt,
    ) = load_docx()

    input_path = Path(args.input_json).expanduser().resolve()
    output_path = Path(args.output_docx).expanduser().resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    payload = json.loads(input_path.read_text(encoding="utf-8-sig"))
    base_dir = input_path.parent
    reference_targets = collect_reference_targets(payload)
    bookmark_state = {"next_id": 1}

    document = Document()
    configure_styles(document, WD_STYLE_TYPE, WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, qn, Pt)

    metadata = payload.get("metadata", {})
    if metadata.get("title"):
        document.add_paragraph(str(metadata["title"]), style="Title")
    if metadata.get("authors"):
        add_text_paragraph(
            document,
            str(metadata["authors"]),
            "PaperAuthors",
            reference_targets,
            bookmark_state,
            OxmlElement,
            qn,
        )
    if metadata.get("affiliations"):
        affiliations = metadata["affiliations"]
        if isinstance(affiliations, list):
            for item in affiliations:
                add_text_paragraph(
                    document,
                    str(item),
                    "PaperAffiliations",
                    reference_targets,
                    bookmark_state,
                    OxmlElement,
                    qn,
                )
        else:
            add_text_paragraph(
                document,
                str(affiliations),
                "PaperAffiliations",
                reference_targets,
                bookmark_state,
                OxmlElement,
                qn,
            )
    if metadata.get("doi"):
        add_text_paragraph(
            document,
            str(metadata["doi"]),
            "PaperMetadata",
            reference_targets,
            bookmark_state,
            OxmlElement,
            qn,
        )
    if metadata.get("subtitle"):
        add_text_paragraph(
            document,
            str(metadata["subtitle"]),
            "PaperMetadata",
            reference_targets,
            bookmark_state,
            OxmlElement,
            qn,
        )

    if "blocks" in payload:
        for block in payload.get("blocks", []):
            add_block(
                document,
                Inches,
                WD_ALIGN_PARAGRAPH,
                block,
                base_dir,
                reference_targets,
                bookmark_state,
                OxmlElement,
                parse_xml,
                qn,
            )
    elif "segments" in payload:
        segments = payload.get("segments", [])
        validate_workbook_segments(segments)
        for segment in segments:
            translated_text = (segment.get("translated_text") or "").strip()
            segment_type = segment.get("type")
            if segment_type != "equation" and not translated_text:
                continue

            if segment_type == "caption":
                block = {"type": "caption", "text": translated_text}
            elif segment_type == "heading":
                block = {"type": "heading", "level": 1, "text": translated_text}
            elif segment_type == "metadata":
                block = {"type": "metadata", "text": translated_text}
            elif segment_type == "equation":
                block = {
                    "type": "equation",
                    "text": translated_text,
                    "equation_omml": segment.get("equation_omml"),
                    "source_id": segment.get("source_id"),
                }
            elif segment_type == "reference":
                block = {"type": "reference", "text": translated_text, "source_id": segment.get("source_id")}
            else:
                block = {"type": "paragraph", "text": translated_text, "source_id": segment.get("source_id")}

            if segment_type in {"caption", "heading", "metadata"}:
                block["source_id"] = segment.get("source_id")

            add_block(
                document,
                Inches,
                WD_ALIGN_PARAGRAPH,
                block,
                base_dir,
                reference_targets,
                bookmark_state,
                OxmlElement,
                parse_xml,
                qn,
            )

    document.save(str(output_path))
    print(str(output_path))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
