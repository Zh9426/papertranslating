#!/usr/bin/env python3
"""Validate a translated DOCX against quality gates for paper-translating."""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path
from zipfile import ZipFile


MOJIBAKE_MARKERS = (
    "锟",
    "鈥",
    "銆",
    "閿",
    "闂",
    "鏉",
)

DISPLAY_FORMULA_RE = re.compile(
    r"(?:\b(?:NMSE|PSNR|SNR|Correlation|VDR|MSE)\b|[=∑Σϕφψλμωαβγδρτ]|\|\||\^|_2|\b[a-zA-Z]\s*=\s*[-+0-9(])"
)
DOC_CITATION_RE = re.compile(r"\[(\d+(?:\s*[,;]\s*\d+)*)\]")
REFERENCE_BOOKMARK_RE = re.compile(r'w:bookmarkStart[^>]+w:name="(ref-\d+)"')
INTERNAL_HYPERLINK_RE = re.compile(r'w:hyperlink[^>]+w:anchor="([^"]+)"')
EQUATION_REF_RE = re.compile(r"\bEq\.?\s*\(?\d+\)?", re.IGNORECASE)
SEGMENT_BOOKMARK_RE = re.compile(r'w:bookmarkStart[^>]+w:name="(seg-[^"]+)"')
SOURCE_EQUATION_NUMBER_RE = re.compile(r"\((\d+)\)")
FIGURE_REF_RE = re.compile(r"\b(?:Fig\.?|Figure)\s*\d+[a-z]?", re.IGNORECASE)
SUPPLEMENTARY_REF_RE = re.compile(
    r"\bSupplementary\s+(?:Fig(?:ure)?\.?|Table|Movie)\s*\d+[a-z]?",
    re.IGNORECASE,
)
CAPTION_PREFIX_RE = re.compile(r"^(?:fig\.?|figure|图)\s*\.?\s*(\d+[a-z]?)", re.IGNORECASE)
PARAMETER_SYMBOL_RE = re.compile(
    r"\b(P|OD|DC|SNR|PSNR|NMSE|Correlation|VDR|f\s*0)\s*=",
    re.IGNORECASE,
)


def load_docx():
    from docx import Document  # type: ignore

    return Document


def load_segments(path: Path) -> list[dict]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    return payload.get("segments", [])


def count_source_figures(manifest_path: Path) -> int:
    payload = json.loads(manifest_path.read_text(encoding="utf-8"))
    return sum(len(page.get("figures", [])) for page in payload.get("pages", []))


def normalize_token(value: str) -> str:
    return re.sub(r"\s+", "", value).lower()


def normalize_markers(markers: set[str]) -> set[str]:
    return {normalize_token(marker) for marker in markers}


def segment_bookmark_name(source_id: str) -> str:
    normalized = re.sub(r"[^A-Za-z0-9_-]+", "-", (source_id or "").strip()).strip("-").lower()
    if not normalized:
        raise ValueError("source_id is required for segment bookmarks")
    if normalized[0].isdigit():
        normalized = f"s-{normalized}"
    return f"seg-{normalized}"


def paragraph_style_counter(document) -> Counter:
    counts: Counter = Counter()
    for para in document.paragraphs:
        text = para.text.strip()
        if text:
            counts[para.style.name] += 1
    return counts


def find_caption_style_mismatches(document) -> list[str]:
    mismatches: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        lowered = text.lower()
        if lowered.startswith("fig.") or lowered.startswith("figure ") or text.startswith("图"):
            if para.style.name != "Caption":
                mismatches.append(text[:160])
    return mismatches


def find_mojibake_paragraphs(document) -> list[str]:
    bad: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        hits = sum(text.count(marker) for marker in MOJIBAKE_MARKERS)
        if hits >= 2:
            bad.append(text[:160])
    return bad


def paragraph_has_omml(paragraph) -> bool:
    xml = paragraph._p.xml
    return "m:oMath" in xml or "m:oMathPara" in xml


def paragraph_has_drawing(paragraph) -> bool:
    return "w:drawing" in paragraph._p.xml


def find_formula_style_mismatches(document) -> list[str]:
    mismatches: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if DISPLAY_FORMULA_RE.search(text) and para.style.name != "PaperEquation":
            mismatches.append(text[:160])
    return mismatches


def find_plaintext_formula_paragraphs(document) -> list[str]:
    bad: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text or para.style.name != "PaperEquation":
            continue
        if DISPLAY_FORMULA_RE.search(text) and not paragraph_has_omml(para):
            bad.append(text[:160])
    return bad


def find_equation_image_paragraphs(document) -> list[str]:
    bad: list[str] = []
    for para in document.paragraphs:
        if para.style.name != "PaperEquation":
            continue
        if paragraph_has_drawing(para) and not paragraph_has_omml(para):
            bad.append((para.text or "[equation-image]").strip()[:160] or "[equation-image]")
    return bad


def media_count(docx_path: Path) -> int:
    with ZipFile(docx_path) as zf:
        return sum(1 for name in zf.namelist() if name.startswith("word/media/"))


def read_document_xml(docx_path: Path) -> str:
    with ZipFile(docx_path) as zf:
        return zf.read("word/document.xml").decode("utf-8", errors="ignore")


def citation_labels_in_document(document) -> set[str]:
    labels: set[str] = set()
    for para in document.paragraphs:
        if para.style.name == "PaperReference":
            continue
        for match in DOC_CITATION_RE.finditer(para.text or ""):
            for item in re.split(r"\s*[,;]\s*", match.group(1)):
                if item.isdigit():
                    labels.add(item)
    return labels


def numbered_reference_count(document) -> int:
    count = 0
    for para in document.paragraphs:
        if para.style.name != "PaperReference":
            continue
        text = (para.text or "").strip()
        if re.match(r"^(?:\[\d+\]|\d+[\.\)])", text):
            count += 1
    return count


def equation_reference_count(document) -> int:
    count = 0
    for para in document.paragraphs:
        count += len(EQUATION_REF_RE.findall(para.text or ""))
    return count


def collect_source_equation_numbers(segments: list[dict]) -> list[int]:
    numbers: set[int] = set()
    for segment in segments:
        text = segment.get("source_text", "") or ""
        if segment.get("type") == "equation":
            for match in SOURCE_EQUATION_NUMBER_RE.findall(text):
                numbers.add(int(match))
        for match in EQUATION_REF_RE.findall(text):
            digits = re.findall(r"\d+", match)
            numbers.update(int(item) for item in digits)
    return sorted(numbers)


def collect_docx_equation_numbers(document) -> list[int]:
    numbers: set[int] = set()
    for para in document.paragraphs:
        for match in EQUATION_REF_RE.findall(para.text or ""):
            digits = re.findall(r"\d+", match)
            numbers.update(int(item) for item in digits)
    return sorted(numbers)


def caption_identifier(text: str) -> str | None:
    match = CAPTION_PREFIX_RE.match((text or "").strip())
    if not match:
        return None
    return f"fig-{normalize_token(match.group(1))}"


def extract_parameter_symbols(text: str) -> set[str]:
    return {normalize_token(match.group(1)) for match in PARAMETER_SYMBOL_RE.finditer(text or "")}


def compare_caption_parameters(source_segments: list[dict], document) -> list[dict]:
    source_expectations: dict[str, set[str]] = {}
    for segment in source_segments:
        if segment.get("type") != "caption":
            continue
        identifier = segment.get("label") or caption_identifier(segment.get("source_text", ""))
        if not identifier:
            continue
        symbols = extract_parameter_symbols(segment.get("source_text", ""))
        if symbols:
            source_expectations[normalize_token(identifier)] = symbols

    final_captions: dict[str, set[str]] = {}
    for para in document.paragraphs:
        if para.style.name != "Caption":
            continue
        identifier = caption_identifier(para.text or "")
        if not identifier:
            continue
        final_captions[normalize_token(identifier)] = extract_parameter_symbols(para.text or "")

    missing: list[dict] = []
    for identifier, expected in sorted(source_expectations.items()):
        absent = sorted(expected - final_captions.get(identifier, set()))
        if absent:
            missing.append(
                {
                    "caption_identifier": identifier,
                    "missing_parameter_symbols": absent,
                }
            )
    return missing


def extract_docx_structural_markers(document) -> dict[str, set[str]]:
    joined = "\n".join((para.text or "").strip() for para in document.paragraphs if (para.text or "").strip())
    return {
        "figure_refs": normalize_markers(set(FIGURE_REF_RE.findall(joined))),
        "equation_refs": normalize_markers(set(EQUATION_REF_RE.findall(joined))),
        "supplementary_refs": normalize_markers(set(SUPPLEMENTARY_REF_RE.findall(joined))),
    }


def compare_source_markers_to_docx(segments: list[dict], document) -> list[dict]:
    source_text = "\n".join(
        (segment.get("source_text") or "").strip() for segment in segments if segment.get("source_text")
    )
    source_markers = {
        "figure_refs": normalize_markers(set(FIGURE_REF_RE.findall(source_text))),
        "equation_refs": normalize_markers(set(EQUATION_REF_RE.findall(source_text))),
        "supplementary_refs": normalize_markers(set(SUPPLEMENTARY_REF_RE.findall(source_text))),
    }
    docx_markers = extract_docx_structural_markers(document)

    missing: list[dict] = []
    for key in ("figure_refs", "equation_refs", "supplementary_refs"):
        absent = sorted(source_markers[key] - docx_markers[key])
        if absent:
            missing.append({"marker_type": key, "missing_values": absent})
    return missing


def compare_source_ids_to_docx(segments: list[dict], document_xml: str) -> list[str]:
    expected = [segment_bookmark_name(str(segment.get("source_id", ""))) for segment in segments if segment.get("source_id")]
    present = set(SEGMENT_BOOKMARK_RE.findall(document_xml))
    return [segment.get("source_id") for segment in segments if segment.get("source_id") and segment_bookmark_name(str(segment["source_id"])) not in present]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", help="Translated DOCX to validate")
    parser.add_argument("--manifest", help="Optional source manifest.json from extract_pdf_assets.py")
    parser.add_argument("--segments", help="Optional source segments JSON from prepare_translation_segments.py")
    args = parser.parse_args()

    Document = load_docx()
    docx_path = Path(args.docx).expanduser().resolve()
    document = Document(str(docx_path))
    document_xml = read_document_xml(docx_path)

    reference_bookmarks = set(REFERENCE_BOOKMARK_RE.findall(document_xml))
    hyperlink_anchors = set(INTERNAL_HYPERLINK_RE.findall(document_xml))
    cited_labels = citation_labels_in_document(document)
    numbered_references = numbered_reference_count(document)
    equation_refs = equation_reference_count(document)
    missing_reference_targets = sorted(
        f"ref-{label}" for label in cited_labels if f"ref-{label}" not in hyperlink_anchors
    )

    report = {
        "docx": str(docx_path),
        "media_count": media_count(docx_path),
        "style_counts": dict(paragraph_style_counter(document)),
        "caption_style_mismatches": find_caption_style_mismatches(document),
        "formula_style_mismatches": find_formula_style_mismatches(document),
        "plaintext_formula_paragraphs": find_plaintext_formula_paragraphs(document),
        "equation_image_paragraphs": find_equation_image_paragraphs(document),
        "equation_reference_count": equation_refs,
        "numbered_reference_count": numbered_references,
        "reference_bookmark_count": len(reference_bookmarks),
        "citation_hyperlink_count": len(hyperlink_anchors),
        "missing_reference_targets": missing_reference_targets,
        "mojibake_paragraphs": find_mojibake_paragraphs(document),
        "missing_equation_numbers": [],
        "missing_equation_bodies": [],
        "missing_caption_parameters": [],
        "missing_structural_markers_from_docx": [],
        "missing_source_ids_from_docx": [],
    }

    if args.manifest:
        manifest_path = Path(args.manifest).expanduser().resolve()
        report["source_figure_count"] = count_source_figures(manifest_path)
        report["missing_figure_count"] = max(0, report["source_figure_count"] - report["media_count"])
        report["suspected_equation_images"] = (
            report["media_count"] > report["source_figure_count"]
            and report["equation_reference_count"] > 0
            and report["style_counts"].get("PaperEquation", 0) == 0
        )
    else:
        report["suspected_equation_images"] = False

    if args.segments:
        segments_path = Path(args.segments).expanduser().resolve()
        segments = load_segments(segments_path)
        source_numbers = collect_source_equation_numbers(segments)
        final_numbers = collect_docx_equation_numbers(document)
        report["source_equation_numbers"] = source_numbers
        report["docx_equation_numbers"] = final_numbers
        report["missing_equation_numbers"] = [number for number in source_numbers if number not in final_numbers]
        report["missing_caption_parameters"] = compare_caption_parameters(segments, document)
        report["missing_structural_markers_from_docx"] = compare_source_markers_to_docx(segments, document)
        report["missing_source_ids_from_docx"] = compare_source_ids_to_docx(segments, document_xml)

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))

    blocking_issues = []
    if report["caption_style_mismatches"]:
        blocking_issues.append("caption_style_mismatches")
    if report["formula_style_mismatches"]:
        blocking_issues.append("formula_style_mismatches")
    if report["plaintext_formula_paragraphs"]:
        blocking_issues.append("plaintext_formula_paragraphs")
    if report["equation_image_paragraphs"]:
        blocking_issues.append("equation_image_paragraphs")
    if report["suspected_equation_images"]:
        blocking_issues.append("suspected_equation_images")
    if report["numbered_reference_count"] > 0 and report["reference_bookmark_count"] == 0:
        blocking_issues.append("missing_reference_bookmarks")
    if report["missing_reference_targets"]:
        blocking_issues.append("missing_reference_targets")
    if report["mojibake_paragraphs"]:
        blocking_issues.append("mojibake_paragraphs")
    if report.get("missing_figure_count", 0) > 0:
        blocking_issues.append("missing_figure_count")
    if report["missing_equation_numbers"]:
        blocking_issues.append("missing_equation_numbers")
    if report["missing_caption_parameters"]:
        blocking_issues.append("missing_caption_parameters")
    if report["missing_structural_markers_from_docx"]:
        blocking_issues.append("missing_structural_markers_from_docx")
    if report["missing_source_ids_from_docx"]:
        blocking_issues.append("missing_source_ids_from_docx")

    return 2 if blocking_issues else 0


if __name__ == "__main__":
    raise SystemExit(main())
